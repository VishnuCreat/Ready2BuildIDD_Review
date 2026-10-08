import csv
import importlib.util
import logging
from pathlib import Path

log = logging.getLogger(__name__)
SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".txt"}


def available_extensions():
    """Formats readable with the currently installed document libraries."""
    result = {".csv", ".txt"}
    if importlib.util.find_spec("openpyxl"):
        result.add(".xlsx")
    if importlib.util.find_spec("pypdf"):
        result.add(".pdf")
    if importlib.util.find_spec("docx"):
        result.add(".docx")
    return result


def missing_document_readers():
    """Return install hints for formats whose optional reader is unavailable."""
    return {
        suffix: package for suffix, module, package in (
            (".pdf", "pypdf", "pypdf"),
            (".docx", "docx", "python-docx"),
            (".xlsx", "openpyxl", "openpyxl"),
        ) if importlib.util.find_spec(module) is None
    }


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported attachment format: {suffix or 'unknown'}")
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="", errors="replace") as f:
            return "\n".join(", ".join(row) for row in csv.reader(f))
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install the documents extra to read PDF files") from exc
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install the documents extra to read DOCX files") from exc
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install the documents extra to read XLSX files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    return "\n".join(" | ".join(str(c) if c is not None else "" for c in row) for sheet in workbook for row in sheet.iter_rows(values_only=True))
